import openpyxl
import pytest
import xlwt

import extractors
from helpers import make_pdf, make_zip


def test_normalize_casefold_and_nfc():
    assert extractors.normalize("ÉCOLE", False) == "école"
    assert extractors.normalize("ÉCOLE", True) == "ÉCOLE"


def test_xlsx_cells(tmp_path):
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["A2"] = "Prix XYZ", 12.0, "autre"
    wb.save(p)
    text = extractors.native_text(str(p))
    assert "Prix XYZ\t12" in text and "autre" in text


def test_xls_cells(tmp_path):
    p = tmp_path / "a.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S1")
    ws.write(0, 0, "Lot XYZ")
    ws.write(0, 1, 3.0)
    wb.save(str(p))
    assert "Lot XYZ\t3" in extractors.native_text(str(p))


def test_docx_runs_joined_and_paragraphs_split(tmp_path):
    xml = ('<w:document><w:body><w:p><w:r><w:t>Expression X</w:t></w:r><w:r><w:t>YZ ici</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>Suite</w:t></w:r></w:p></w:body></w:document>')
    p = make_zip(tmp_path / "a.docx", {"word/document.xml": xml, "[Content_Types].xml": "<x/>"})
    text = extractors.native_text(p)
    assert "Expression XYZ ici" in text and "ici\nSuite" in text


def test_pptx_slides(tmp_path):
    p = make_zip(tmp_path / "a.pptx", {"ppt/slides/slide1.xml": "<p:sld><a:p><a:r><a:t>Diapo XYZ</a:t></a:r></a:p></p:sld>"})
    assert "Diapo XYZ" in extractors.native_text(p)


def test_odt_content(tmp_path):
    p = make_zip(tmp_path / "a.odt", {"content.xml": "<office:text><text:p>Texte<text:s/>XYZ &amp; co</text:p></office:text>"})
    assert "Texte XYZ & co" in extractors.native_text(p)


def test_pdf_text(tmp_path):
    p = tmp_path / "a.pdf"
    make_pdf(p, "Bonjour XYZ")
    assert "Bonjour XYZ" in extractors.native_text(str(p))


@pytest.mark.parametrize("data", [
    "Café XYZ".encode("utf-8"),
    "Café XYZ".encode("cp1252"),
    b"\xff\xfe" + "Café XYZ".encode("utf-16-le"),
    b"\xef\xbb\xbf" + "Café XYZ".encode("utf-8"),
])
def test_text_encodings(tmp_path, data):
    p = tmp_path / "a.txt"
    p.write_bytes(data)
    assert extractors.native_text(str(p)) == "Café XYZ"


def test_unknown_extension_returns_none(tmp_path):
    p = tmp_path / "a.zzz"
    p.write_bytes(b"XYZ")
    assert extractors.native_text(str(p)) is None


def test_encrypted_ooxml_is_protected(tmp_path):
    p = tmp_path / "a.xlsx"
    p.write_bytes(extractors.OLE_MAGIC + b"\x00" * 600)
    with pytest.raises(extractors.ProtectedFile):
        extractors.native_text(str(p))


def test_text_limit_applied_while_reading_xlsx(tmp_path, monkeypatch):
    p = tmp_path / "big.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(2000):
        ws.append([f"ligne numéro {i}", "x" * 20])
    wb.save(p)
    monkeypatch.setattr(extractors.config, "TEXT_LIMIT", 500)
    text = extractors.native_text(str(p))
    assert 500 < len(text) < 700


def test_text_limit_applied_while_reading_pptx(tmp_path, monkeypatch):
    members = {
        f"ppt/slides/slide{i}.xml": f"<p:sld><a:p>diapo {i} " + "y" * 80 + "</a:p></p:sld>"
        for i in range(100)
    }
    p = make_zip(tmp_path / "big.pptx", members)
    monkeypatch.setattr(extractors.config, "TEXT_LIMIT", 500)
    text = extractors.native_text(p)
    assert 0 < len(text) <= 500
